import openpyxl
from pathlib import Path

WB = Path(r"H:\.shortcut-targets-by-id\15cPT84Tcymc9b2jqyLRYJjcuEXIBOfOp\0. Business\Multi Family\!!! W.C. Portfolio\!!! Brio\Financial\202605\Brio_Financial_Analysis_202605.xlsx")
wb = openpyxl.load_workbook(WB, data_only=True)

def num(v):
    return v if isinstance(v, (int, float)) else 0.0

# ---------------- Phase 3: below-the-line (T12 rows after NOI 221) ----------------
t12 = wb["T12"]
REVIEW = 14
M3 = [11, 12, 13]
print("=" * 90)
print("PHASE 3 — BELOW-THE-LINE (after NOI, GL 70000+)")
print("=" * 90)
below_total_rev = 0.0
for r in range(222, 270):
    code = t12.cell(r, 1).value
    name = t12.cell(r, 2).value
    if not name:
        continue
    rev = num(t12.cell(r, REVIEW).value)
    total = num(t12.cell(r, 15).value)
    a3 = sum(num(t12.cell(r, c).value) for c in M3) / 3
    is_sub = isinstance(code, str) and code.endswith(("-099", "-090", "-199", "-999", "-098"))
    if is_sub or name.isupper() or "Total" in name:
        if name.strip() in ("Total Routine Replacement Expense", "Total Capital / Renovation Expense",
                             "Total Debt Service", "Net Income", "NOI After Replacements",
                             "Total Partnership / Owner Expenses"):
            print(f"  [SUBTOTAL] {code} {name.strip():<45} rev={rev:>13,.0f} 12mo-total={total:>14,.0f}")
        continue
    if rev == 0 and total == 0:
        continue
    flag = ""
    if (abs(rev) > 10000 and a3 != 0 and abs(rev) > 2 * abs(a3)) or abs(rev) > 25000:
        flag = "  <<< FLAG (large/unusual)"
    print(f"    {code} {name.strip():<45} rev={rev:>13,.0f} 3mo={a3:>11,.0f} 12mo-total={total:>13,.0f}{flag}")

# NOI bridge
noi = num(t12.cell(221, REVIEW).value)
net = num(t12.cell(269, REVIEW).value)
print(f"\n  NOI (May) = {noi:,.0f}   ->   Net Income = {net:,.0f}   (below-line drag = {net-noi:,.0f})")
debt = [num(t12.cell(263, c).value) for c in range(3, 15)]
print(f"  Debt service monthly (Jun..May): {[f'{x:,.0f}' for x in debt]}")

# ---------------- Phase 4: Budget Comparison ----------------
bc = wb["Budget Comparison"]
# cols: A name, B PTDact, C PTDbud, D PTDvar, E PTD%, F YTDact, G YTDbud, H YTDvar, I YTD%, J annual
items = []
for r in range(6, 252):
    name = bc.cell(r, 1).value
    if not name or "Total" in name or name.strip().isupper():
        continue
    pa, pb, pv, pp = (num(bc.cell(r, c).value) for c in (2, 3, 4, 5))
    ya, yb, yv, yp = (num(bc.cell(r, c).value) for c in (6, 7, 8, 9))
    ann = num(bc.cell(r, 10).value)
    if pa == 0 and pb == 0 and ya == 0 and yb == 0:
        continue
    items.append(dict(r=r, name=name.strip(), pa=pa, pb=pb, pv=pv, pp=pp,
                      ya=ya, yb=yb, yv=yv, yp=yp, ann=ann))

def show(title, key, n=8):
    print(f"\n--- {title} ---")
    for it in sorted(items, key=key)[:n]:
        print(f"  {it['name']:<42} PTD a/b={it['pa']:>10,.0f}/{it['pb']:>10,.0f} var={it['pv']:>10,.0f} ({it['pp']:>6.1f}%)"
              f"  | YTD var={it['yv']:>11,.0f} ({it['yp']:>6.1f}%)")

print("\n" + "=" * 90)
print("PHASE 4 — BUDGET COMPARISON")
print("=" * 90)
show("Top PTD UNFAVORABLE by $ (most negative variance)", lambda x: x["pv"])
show("Top PTD FAVORABLE by $", lambda x: -x["pv"])
show("Top YTD UNFAVORABLE by $ (most negative variance)", lambda x: x["yv"])
show("Top YTD FAVORABLE by $", lambda x: -x["yv"])
show("Top PTD UNFAVORABLE by % (material >$1k)", lambda x: x["pp"] if abs(x["pv"]) > 1000 else 999)
show("Top YTD UNFAVORABLE by % (material >$3k)", lambda x: x["yp"] if abs(x["yv"]) > 3000 else 999)

# pacing: 5 months elapsed (Jan-May 2026 fiscal? budget YTD covers Jan-May = 5 months)
print("\n--- Budget threshold flags ---")
for it in items:
    fl = []
    if it["pv"] < 0 and abs(it["pp"]) > 50 and abs(it["pv"]) > 2500:
        fl.append("PTD High")
    elif it["pv"] < 0 and abs(it["pp"]) > 20 and abs(it["pv"]) > 1000:
        fl.append("PTD Mod")
    if it["yv"] < 0 and abs(it["yp"]) > 30 and abs(it["yv"]) > 10000:
        fl.append("YTD High")
    elif it["yv"] < 0 and abs(it["yp"]) > 15 and abs(it["yv"]) > 5000:
        fl.append("YTD Mod")
    if fl:
        print(f"  [{','.join(fl):16}] {it['name']:<42} PTDvar={it['pv']:>10,.0f} YTDvar={it['yv']:>11,.0f}")
