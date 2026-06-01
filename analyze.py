"""Phase 2-4 analysis. Reconstructed dollar thresholds (skill text arrived with
corrupted $ values); using standard materiality cutoffs noted in comments."""
import openpyxl
from pathlib import Path

WB = Path(r"H:\.shortcut-targets-by-id\15cPT84Tcymc9b2jqyLRYJjcuEXIBOfOp\0. Business\Multi Family\!!! W.C. Portfolio\!!! Brio\Financial\202605\Brio_Financial_Analysis_202605.xlsx")
wb = openpyxl.load_workbook(WB, data_only=True)
t12 = wb["T12"]

# T12 month columns: C(3)=Jun25 ... N(14)=May26 ; review=14, prior=13, 3mo=11,12,13
REVIEW, PRIOR = 14, 13
M3 = [11, 12, 13]
M12 = list(range(3, 15))
NOI_ROW = 221  # 69999-099 Net Operating Income

def num(v):
    return v if isinstance(v, (int, float)) else 0.0

def avg(vals):
    return sum(vals) / len(vals) if vals else 0.0

rows = []
for r in range(6, NOI_ROW):
    code = t12.cell(r, 1).value
    name = t12.cell(r, 2).value
    if not name:
        continue
    months = [num(t12.cell(r, c).value) for c in M12]
    rev = num(t12.cell(r, REVIEW).value)
    prior = num(t12.cell(r, PRIOR).value)
    a3 = avg([num(t12.cell(r, c).value) for c in M3])
    a12 = avg(months)
    total = num(t12.cell(r, 15).value)
    rows.append(dict(r=r, code=code, name=name.strip(), rev=rev, prior=prior,
                     a3=a3, a12=a12, total=total, months=months))

def is_subtotal(code, name):
    if code and isinstance(code, str):
        for suf in ("-099", "-098", "-090", "-999", "-199"):
            if code.endswith(suf):
                return True
    if "Total" in name or name.isupper():
        return True
    return False

print("=" * 100)
print("PHASE 2 — ABOVE-NOI FLAGS (review month = May 2026, col N)")
print("=" * 100)

flags = []
for d in rows:
    code, name = d["code"], d["name"]
    if is_subtotal(code, name):
        continue
    rev, prior, a3, a12 = d["rev"], d["prior"], d["a3"], d["a12"]
    # materiality skip (except negatives handled below)
    gl = str(code) if code else ""
    is_income = gl[:1] == "4"
    is_opex = gl[:1] in ("5", "6")
    sev = None
    reasons = []

    # negative flags (never skipped on materiality)
    primary_rent = name in ("Market Rent", "Gain / Loss To Lease")
    contra = any(k in name for k in ("Concession", "Vacancy Loss", "Bad Debt", "Loss To Lease",
                                     "Employee Units", "Model & Storage"))
    if is_income and rev < 0 and not contra and not primary_rent:
        sev = "High"; reasons.append(f"NEGATIVE other rental income {rev:,.0f}")
    if is_opex and rev < 0:
        sev = "High"; reasons.append(f"NEGATIVE expense {rev:,.0f}")

    immaterial = abs(rev) < 500 and abs(a3) < 500
    if immaterial and not reasons:
        continue

    pm_swing = (rev - prior) / abs(prior) if prior else None
    a3_swing = (rev - a3) / abs(a3) if a3 else None

    if is_income and not contra:
        if pm_swing is not None and abs(pm_swing) > 0.20 and abs(rev - prior) > 1000:
            reasons.append(f"vs prior {pm_swing:+.0%} ({rev:,.0f} vs {prior:,.0f})"); sev = sev or "Moderate"
        if a3_swing is not None and abs(a3_swing) > 0.25 and abs(rev - a3) > 2500:
            reasons.append(f"vs 3mo-avg {a3_swing:+.0%} ({rev:,.0f} vs {a3:,.0f})"); sev = sev or "Moderate"
        if rev == 0 and a3 > 0:
            reasons.append("revenue dropped to ZERO"); sev = "High"
    if "One-Time Concessions" in name and a3 != 0 and abs(rev) > 1.5 * abs(a3):
        reasons.append(f"One-Time Concessions {abs(rev)/abs(a3):.0%} of 3mo-avg"); sev = sev or "Moderate"
    if "Bad Debt - Rent" == name and a3 != 0 and abs(rev) > 2 * abs(a3) and abs(rev) > 5000:
        reasons.append(f"Bad Debt {abs(rev)/abs(a3):.0%} of 3mo-avg"); sev = "High"
    if "Vacancy Loss" in name and a3 != 0 and abs(rev) > 1.2 * abs(a3) and abs(rev) > 5000:
        reasons.append(f"Vacancy Loss {abs(rev)/abs(a3):.0%} of 3mo-avg"); sev = sev or "Moderate"

    if is_opex:
        if pm_swing is not None and abs(pm_swing) > 0.25 and abs(rev - prior) > 1000:
            reasons.append(f"vs prior {pm_swing:+.0%} ({rev:,.0f} vs {prior:,.0f})"); sev = sev or "Moderate"
        if a3_swing is not None and abs(a3_swing) > 0.30 and abs(rev - a3) > 2500:
            reasons.append(f"vs 3mo-avg {a3_swing:+.0%} ({rev:,.0f} vs {a3:,.0f})"); sev = sev or "Moderate"
        if rev == 0 and a3 > 0:
            reasons.append("expense dropped to zero (3mo-avg>0)"); sev = sev or "Moderate"
        if d["total"] and abs(rev) > 0.5 * abs(d["total"]):
            reasons.append(f"single month = {rev/d['total']:.0%} of 12mo total"); sev = "High"
        if name == "Ad Valorem Property Taxes" and prior and abs(rev - prior) / abs(prior) > 0.10:
            reasons.append(f"Ad Valorem differs {(rev-prior)/prior:+.0%} from prior"); sev = sev or "Moderate"
        if "Property Insurance" in name and prior and abs(rev - prior) > 1:
            reasons.append(f"Insurance changed from prior ({rev:,.0f} vs {prior:,.0f})"); sev = sev or "Moderate"
        utility_names = ("Electric", "Water", "Gas", "Utility")
        if any(u in name for u in utility_names) and a3 != 0 and abs(rev) > 1.5 * abs(a3) and abs(rev) > 2000:
            reasons.append(f"utility {abs(rev)/abs(a3):.0%} of 3mo-avg"); sev = sev or "Moderate"

    if reasons:
        flags.append(dict(d=d, sev=sev or "Moderate", reasons=reasons))

flags.sort(key=lambda f: (0 if f["sev"] == "High" else 1, -abs(f["d"]["rev"])))
for f in flags:
    d = f["d"]
    print(f"[{f['sev']:8}] {d['code']} {d['name']:<42} rev={d['rev']:>12,.0f} prior={d['prior']:>11,.0f} 3mo={d['a3']:>11,.0f}")
    for rsn in f["reasons"]:
        print(f"            - {rsn}")

# Management fee rate check
tot_income = num(t12.cell(58, REVIEW).value)
mgmt = num(t12.cell(205, REVIEW).value)
print(f"\nMgmt fee rate = {mgmt:,.0f} / {tot_income:,.0f} = {mgmt/tot_income:.2%} (target 2.5%-4.0%)")
print(f"\nTotal flags: {len(flags)}  (High={sum(1 for f in flags if f['sev']=='High')})")
