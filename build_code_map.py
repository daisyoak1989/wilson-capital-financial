"""One-off: build a union account-name -> GL-code map across every portfolio
T12 that carries the GL-code column, so combine.py can backfill the column for
exports (e.g. Local DS) that omit it. The chart of accounts (Appfolio ysi_is
tree) is standardized portfolio-wide."""
import openpyxl, re, glob, os, json

BASE = r"H:\.shortcut-targets-by-id\15cPT84Tcymc9b2jqyLRYJjcuEXIBOfOp\0. Business\Multi Family\!!! W.C. Portfolio"
OUT = os.path.join(os.path.dirname(__file__), "account_codes.json")

def norm(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower() if s else ""

files = glob.glob(os.path.join(BASE, "*", "Financial", "**", "12_Month_Statement_*.xlsx"), recursive=True)
print("found", len(files), "T12 files")
union = {}
for f in files:
    try:
        ws = openpyxl.load_workbook(f, data_only=True).active
    except Exception as e:
        print("ERR", f, e)
        continue
    cnt = 0
    for r in range(6, ws.max_row + 1):
        c = ws.cell(r, 1).value
        n = ws.cell(r, 2).value
        if isinstance(c, str) and re.match(r"^\d{5}-\d{3}$", c.strip()) and n:
            union.setdefault(norm(n), c.strip())
            cnt += 1
    if cnt:
        prop = os.path.basename(os.path.dirname(os.path.dirname(f)))
        print(f"  {prop:30s} {os.path.basename(f):45s} codes:{cnt}")
print("UNION size:", len(union))
with open(OUT, "w", encoding="utf-8") as fh:
    json.dump(union, fh, indent=0, ensure_ascii=False)
print("Saved:", OUT)
