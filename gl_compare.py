"""GL cross-month comparison (house rules C.8 & C.9).

The GL export only ever contains the *review month's* transactions, so to judge
whether a recurring contract/subscription moved for a real reason you must open
the PREVIOUS month's folder and read the same GL section there. This script does
that automatically.

Usage (run from the repo with venv\\Scripts\\python.exe -X utf8):
  gl_compare.py --wb "<current combined workbook>" --auto
      C.8: dump every NEGATIVE expense's current-month GL (no dollar floor).
      C.9: for every stable recurring line that moved materially, show the
           current vs previous-month GL side-by-side.
  gl_compare.py --wb "<current workbook>" 53180-000 58090-000
      Compare the given code(s) current vs previous month.
  gl_compare.py --wb "<current workbook>" --stable     # only the C.9 movers
  gl_compare.py --wb "<current workbook>" --neg        # only the C.8 negatives

The previous month folder is resolved by decrementing the YYYYMM folder name
that contains the current workbook (so Bower's extra nesting is handled too).
"""
import sys
import re
import glob
import openpyxl
from pathlib import Path
from datetime import datetime, timedelta

# Recurring lines that should be roughly level month-to-month. Deliberately
# scoped to contracts / subscriptions / utilities-style services — NOT lumpy
# R&M supplies or make-ready contractors.
STABLE_KEYS = (
    "Contract", "Subscription", "Software", "Licenses", "Automation",
    "Internet Access", "Telephone", "Cellular", "Answering", "Music",
    "Resident Screening", "Copy Machine", "Bank Charges", "Trash",
    "Landscape", "Elevator", "Cable", "Pest Control",
    "Access Gate", "Alarm", "Janitorial", "Pool", "Security",
)
# NB: "Insurance" is intentionally excluded — Property Insurance is self-checked
# against the portfolio reference table (A.5) and Group Insurance is a payroll
# benefit (folded into the payroll cluster), so neither belongs in the C.9 sweep.
# Word-boundary match so "Contract" does NOT catch "Paint Contractor".
_STABLE_RE = re.compile(r"\b(" + "|".join(re.escape(k) for k in STABLE_KEYS) + r")\b", re.I)
# Lumpy lines that happen to brush a keyword but are NOT level recurring spend.
_LUMPY_EXCLUDE = ("Paint", "Housekeeper", "Cleaning Service", "Make", "Supplies")


def is_stable_line(name):
    """True for recurring contract/subscription lines that should be ~level."""
    nm = str(name)
    if any(x in nm for x in _LUMPY_EXCLUDE):
        return False
    return bool(_STABLE_RE.search(nm))


STABLE_PCT = 0.25     # deviation vs 3-mo avg to flag
STABLE_DOLLAR = 100   # ...and a minimum dollar move so tiny lines stay quiet


def num(v):
    return v if isinstance(v, (int, float)) else 0.0


def serial_to_date(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float)) and v > 30000:
        return (datetime(1899, 12, 30) + timedelta(days=v)).strftime("%Y-%m-%d")
    return str(v) if v is not None else ""


def gl_worksheet(path):
    """Return the GL worksheet from a combined workbook (tab 'GL') or a raw
    GeneralLedger export (first sheet)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    return wb["GL"] if "GL" in wb.sheetnames else wb.worksheets[0]


def parse_headers(ws):
    """Section header rows: column A holds a GL code like 54050-000."""
    hdr = []
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and len(a) >= 8 and a[5:6] == "-" and a[:5].isdigit():
            hdr.append((r, a.strip(), ws.cell(r, 5).value))
    return hdr


def dump_section(ws, code, label):
    """Print one GL section's transactions and return its net change."""
    hdr = parse_headers(ws)
    starts = sorted(r for r, _, _ in hdr)
    idx = {c: r for r, c, _ in hdr}
    if code not in idx:
        print(f"   [{label}] {code}: not present (no activity this month)")
        return 0.0
    start = idx[code]
    end = next((r for r in starts if r > start), ws.max_row + 1)
    name = ws.cell(start, 5).value
    print(f"   [{label}] {code}  ({str(name).strip()})")
    print(f"   {'Date':<12}{'Description':<42}{'Debit':>12}{'Credit':>12}")
    deb_t = cred_t = 0.0
    for r in range(start + 1, end):
        date = serial_to_date(ws.cell(r, 3).value)
        desc = str(ws.cell(r, 5).value or "")[:40]
        deb = ws.cell(r, 8).value
        cred = ws.cell(r, 9).value
        if not (date or desc or isinstance(deb, (int, float)) or isinstance(cred, (int, float))):
            continue
        deb_t += num(deb); cred_t += num(cred)
        ds = f"{deb:,.2f}" if isinstance(deb, (int, float)) else ""
        cs = f"{cred:,.2f}" if isinstance(cred, (int, float)) else ""
        print(f"   {date:<12}{desc:<42}{ds:>12}{cs:>12}")
    net = deb_t - cred_t
    print(f"   {'':54}{'  Net = ' + format(net, ',.2f'):>20}")
    return net


def prev_month_folder(wb_path):
    """Sibling YYYYMM folder for the month before the current workbook's folder."""
    month_dir = Path(wb_path).parent
    m = re.match(r"^(\d{4})(\d{2})$", month_dir.name)
    if not m:
        return None
    y, mo = int(m.group(1)), int(m.group(2))
    y, mo = (y - 1, 12) if mo == 1 else (y, mo - 1)
    cand = month_dir.parent / f"{y}{mo:02d}"
    return cand if cand.exists() else None


def find_prev_gl(prev_folder):
    """Prefer a previously-built combined workbook; else the raw GeneralLedger."""
    for pat in ("*Financial_Analysis_*.xlsx", "*GeneralLedger*.xlsx", "*General*Ledger*.xlsx"):
        for p in sorted(prev_folder.glob(pat)):
            if not p.name.startswith("~$"):
                return p
    return None


def scan_t12(wb_path):
    """Return (negatives, movers, rev_label); the lists hold (code, name, rev,
    prior, a3) and rev_label is the review-month header (e.g. 'May 2026')."""
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    if "T12" not in wb.sheetnames:
        return [], [], ""
    ws = wb["T12"]
    mon_re = re.compile(r"^[A-Z][a-z]{2} \d{4}$")
    cols = [c for c in range(1, ws.max_column + 1)
            if isinstance(ws.cell(5, c).value, str) and mon_re.match(ws.cell(5, c).value.strip())]
    if not cols:
        return [], [], ""
    # Baseline = the 3 months BEFORE the review month (the prior run-rate),
    # matching review_build.py's M3 so both tools flag the same lines.
    rev_c, pri_c, m3 = cols[-1], cols[-2], cols[-4:-1]
    rev_label = str(ws.cell(5, rev_c).value or "review mo").strip()
    negatives, movers = [], []
    for r in range(6, ws.max_row + 1):
        code = ws.cell(r, 1).value
        name = ws.cell(r, 2).value
        if not (isinstance(code, str) and re.match(r"^[5678]\d{4}-\d{3}$", code.strip())):
            continue
        if code.strip().endswith(("-099", "-098", "-090", "-199", "-999")):
            continue
        nm = str(name).strip()
        rev = num(ws.cell(r, rev_c).value)
        prior = num(ws.cell(r, pri_c).value)
        a3 = sum(num(ws.cell(r, c).value) for c in m3) / 3
        if rev < 0:
            negatives.append((code.strip(), nm, rev, prior, a3))
        if is_stable_line(nm):
            base = a3 if a3 else prior
            dev = abs(rev - base)
            if base and dev / max(abs(base), 1) > STABLE_PCT and dev > STABLE_DOLLAR:
                movers.append((code.strip(), nm, rev, prior, a3))
    return negatives, movers, rev_label


def main():
    a = sys.argv[1:]
    if "--wb" not in a:
        print(__doc__)
        sys.exit(1)
    i = a.index("--wb")
    wb_path = a[i + 1]
    del a[i:i + 2]
    flags = {x for x in a if x.startswith("--")}
    codes = [x for x in a if not x.startswith("--")]
    auto = "--auto" in flags
    do_neg = auto or "--neg" in flags or (not flags and not codes)
    do_stable = auto or "--stable" in flags or (not flags and not codes)

    cur_gl = gl_worksheet(wb_path)
    prev_folder = prev_month_folder(wb_path)
    prev_gl_path = find_prev_gl(prev_folder) if prev_folder else None
    prev_gl = gl_worksheet(prev_gl_path) if prev_gl_path else None
    cur_label = Path(wb_path).parent.name
    prev_label = prev_folder.name if prev_folder else "prev"
    print(f"Current: {cur_label}    Previous: {prev_label}"
          f"{'' if prev_gl else '  (PREV GL NOT FOUND — current-only)'}")

    # Explicit codes -> straight current-vs-prior comparison.
    if codes:
        for code in codes:
            print("\n" + "=" * 90 + f"\n{code}\n" + "=" * 90)
            dump_section(cur_gl, code, cur_label)
            if prev_gl:
                dump_section(prev_gl, code, prev_label)
        return

    negatives, movers, rev_label = scan_t12(wb_path)

    if do_neg:
        print("\n" + "#" * 90)
        print("# C.8  NEGATIVE EXPENSES (any amount) — current-month GL")
        print("#" * 90)
        if not negatives:
            print("  none.")
        for code, nm, rev, prior, a3 in negatives:
            print(f"\n--- {code} {nm}  (T12 month = {rev:,.2f}) ---")
            dump_section(cur_gl, code, cur_label)

    if do_stable:
        print("\n" + "#" * 90)
        print("# C.9  STABLE RECURRING LINES THAT MOVED — current vs previous-month GL")
        print("#" * 90)
        if not movers:
            print("  none.")
        for code, nm, rev, prior, a3 in movers:
            print(f"\n--- {code} {nm}  ({rev_label} {rev:,.2f} vs prior {prior:,.2f} vs 3-mo {a3:,.2f}) ---")
            dump_section(cur_gl, code, cur_label)
            if prev_gl:
                dump_section(prev_gl, code, prev_label)


if __name__ == "__main__":
    main()
